import xml.etree.ElementTree as ET
from datetime import datetime, timedelta
import pandas as pd
from openpyxl import load_workbook
import shutil
import re
import PyPDF2
from config import *
import os

def extrair_numero_fatura_limpo(texto):
    """
    Extrai um número de fatura do formato 'Nº 000.001.479',
    ignora os blocos com zero, remove zeros à esquerda e remove pontos.
    Exemplo: 'Nº 000.001.479' -> '1479'
    """
    match = re.search(r'Nº\s+((?:\d+\.)*\d+)', texto)
    if match:
        blocos = match.group(1).split(".")
        significativos = [str(int(b)) for b in blocos if int(b) != 0]
        resultado = "".join(significativos)  # Remove pontos usando join sem separador
        return resultado
    return ''

def extrair_texto(caminho_do_pdf):
    texto = ''
    with open(caminho_do_pdf, 'rb') as arquivo:
        leitor_pdf = PyPDF2.PdfReader(arquivo)
        for pagina in leitor_pdf.pages:
            texto_pagina = pagina.extract_text()
            if texto_pagina:
                texto_pagina = texto_pagina.replace('\n', ' ')
                texto += texto_pagina + ' '
    
    if not texto:
        print(f"❌ Erro ao extrair texto do PDF: {caminho_do_pdf}")
    else:
        print(f"✅ Texto extraído do PDF {caminho_do_pdf}: {texto[:500]}...")  # Mostra os primeiros 500 caracteres do texto extraído
    return texto.strip()  # Remove espaços extras no início e no fim

def registro_existe(df, cnpj, data_inicio, data_fim, valor_total):
    return not df[(df['CNPJ'] == cnpj) & (df['DATA INICIO'] == data_inicio) & (df['DATA FIM'] == data_fim) & (df['VALOR TOTAL'] == valor_total)].empty

def todos_campos_preenchidos(informacoes):
    return True

def excluir_arquivo(caminho_arquivo):
    try:
        if os.path.isfile(caminho_arquivo):
            os.remove(caminho_arquivo)
            print(f"🗑️ Arquivo excluído: {caminho_arquivo}")
    except Exception as e:
        print(f"❌ Erro ao excluir arquivo: {e}")

def adicionar_na_planilha(informacoes, caminho_planilha, nome_arquivo):
    try:
        df = pd.read_excel(caminho_planilha)
    except FileNotFoundError:
        print(f"❌ O arquivo '{caminho_planilha}' não foi encontrado. Criando um novo.")
        df = pd.DataFrame(columns=['CNPJ', 'VALOR TOTAL', 'VOLUME TOTAL', 'DATA EMISSAO', 'DATA INICIO', 'DATA FIM', 'NUMERO FATURA', 'VALOR ICMS', 'DISTRIBUIDORA', 'NOME DO ARQUIVO'])
    
    cnpj = informacoes.get('cnpj', '')
    data_inicio = informacoes.get('data_inicio', '')
    data_fim = informacoes.get('data_fim', '')
    valor_total = pd.to_numeric(informacoes.get('valor_total', '0').replace('.', '').replace(',', '.'), errors='coerce')
    volume_total = informacoes.get('volume_total', 0)
    valor_icms = pd.to_numeric(informacoes.get('valor_icms', '0').replace('.', '').replace(',', '.'), errors='coerce')

    if registro_existe(df, cnpj, data_inicio, data_fim, valor_total):
        print(f"❌ Registro duplicado encontrado. Não será inserido.")
        return False 
    
    nova_linha = pd.DataFrame([{
        'CNPJ': cnpj,
        'VALOR TOTAL': valor_total,
        'VOLUME TOTAL': volume_total,
        'DATA EMISSAO': informacoes.get('data_emissao', ''),
        'DATA INICIO': data_inicio,
        'DATA FIM': data_fim,
        'NUMERO FATURA': informacoes.get('numero_fatura', ''),
        'VALOR ICMS': valor_icms,
        'DISTRIBUIDORA': DIST,
        'NOME DO ARQUIVO': nome_arquivo
    }])
    df = pd.concat([df, nova_linha], ignore_index=True)
    df.to_excel(caminho_planilha, index=False)
    print("✅ Dados adicionados com sucesso à planilha.")
    return True

def mover_arquivo(origem, destino):
    shutil.move(origem, destino)
    print(f"✅ Arquivo movido para {destino}")

def verificar_linha_preenchida(caminho_planilha, informacoes):
    try:
        workbook = load_workbook(caminho_planilha)
        sheet = workbook.active

        for row in sheet.iter_rows(min_row=2, values_only=True):  # Ignora o cabeçalho
            if (
                row[0] == informacoes.get('cnpj') and
                row[1] == informacoes.get('valor_total') and
                row[2] == informacoes.get('volume_total') and
                row[3] == informacoes.get('data_emissao') and
                row[4] == informacoes.get('data_inicio') and
                row[5] == informacoes.get('data_fim') and
                row[6] == informacoes.get('numero_fatura') and
                row[7] == informacoes.get('valor_icms')

            ):
                if all(cell is not None and cell != '' for cell in row):
                    return True
                else:
                    return False
        return False  # Retorna False se a linha correspondente não for encontrada
    except Exception as e:
        print(f"❌ Erro ao verificar a planilha: {e}")
        return False

def extrair_informacoes_xml(xml_root):
    data_emissao_iso = xml_root.find('.//nfe:ide/nfe:dhEmi', NAMESPACE).text
    data_emissao_dt = datetime.fromisoformat(data_emissao_iso)

    data_emissao = data_emissao_dt.date()
    data_inicio = data_emissao.replace(day=1)
    next_month = data_emissao.replace(day=28) + timedelta(days=4)
    data_fim = (next_month - timedelta(days=next_month.day))

    informacoes = {
        'cnpj': xml_root.find('.//nfe:dest/nfe:CNPJ', NAMESPACE).text,
        'valor_total': xml_root.find('.//nfe:total/nfe:ICMSTot/nfe:vNF', NAMESPACE).text,
        'volume_total': xml_root.find('.//nfe:det/nfe:prod/nfe:qCom', NAMESPACE).text,
        'data_emissao': data_emissao,
        'data_inicio': data_inicio,
        'data_fim': data_fim,
        'numero_fatura': xml_root.find('.//nfe:ide/nfe:nNF', NAMESPACE).text,
        'valor_icms': xml_root.find('.//nfe:total/nfe:ICMSTot/nfe:vICMS', NAMESPACE).text
    }
    return informacoes

def tratar_datas(informacoes):
    """
    Converte campos de data e CRIA datas padrão se não existirem (para PDFs).
    🔧 AJUSTE: Apenas adiciona data_inicio e data_fim se não existirem.
    """
    from datetime import datetime, date
    from calendar import monthrange
    
    campos_de_data = ['data_emissao', 'data_inicio', 'data_fim']

    # Converter strings para objetos date
    for campo in campos_de_data:
        data_str = informacoes.get(campo)
        if data_str and isinstance(data_str, str):
            try:
                data_str = data_str.replace('-', '/').strip()
                informacoes[campo] = datetime.strptime(data_str, '%d/%m/%Y').date()
                print(f"[DEBUG] {campo} convertida: {informacoes[campo]}")
            except ValueError:
                print(f"⚠️ Erro ao converter '{campo}': formato inválido ({data_str})")
                informacoes[campo] = None

    # 🔧 APENAS para PDFs: Criar data_inicio e data_fim se não existirem
    data_emissao = informacoes.get('data_emissao')
    
    # Verificar se precisa criar as datas (PDF não tem, XML já tem)
    if data_emissao and (informacoes.get('data_inicio') is None or informacoes.get('data_fim') is None):
        print(f"[INFO] PDF detectado - criando datas de fornecimento baseadas na emissão: {data_emissao}")
        
        # Se data_emissao é string, converter
        if isinstance(data_emissao, str):
            try:
                data_emissao = datetime.strptime(data_emissao.replace('-', '/'), '%d/%m/%Y').date()
                informacoes['data_emissao'] = data_emissao
            except ValueError:
                print(f"❌ Erro ao converter data_emissao: {data_emissao}")
                return informacoes
        
        # Extrair ano e mês da data de emissão
        ano = data_emissao.year
        mes = data_emissao.month
        
        # Criar apenas se não existir
        if informacoes.get('data_inicio') is None:
            informacoes['data_inicio'] = date(ano, mes, 1)  # Primeiro dia do mês
            print(f"[INFO] Data início criada: {informacoes['data_inicio']}")
        
        if informacoes.get('data_fim') is None:
            ultimo_dia_do_mes = monthrange(ano, mes)[1]
            informacoes['data_fim'] = date(ano, mes, ultimo_dia_do_mes)  # Último dia do mês
            print(f"[INFO] Data fim criada: {informacoes['data_fim']}")
    
    # Debug final
    print(f"[DEBUG] Datas finais - Emissão: {informacoes.get('data_emissao')}, Início: {informacoes.get('data_inicio')}, Fim: {informacoes.get('data_fim')}")
    
    return informacoes